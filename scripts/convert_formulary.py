"""院内医薬品集（印刷レイアウトのCSV）を 1行=1薬品・1規格 の表に整形する。

  python scripts/convert_formulary.py "~/Downloads/★院内医薬品集 ... .csv"
  python scripts/convert_formulary.py 入力.csv --out data/formulary/院内医薬品集

出力: <out>.csv（UTF-8 BOM付き、Excelでそのまま開ける）と <out>.xlsx

元CSVの構造:
  ≪内用薬≫ などの区分行 → 「112　催眠鎮静剤…」の薬効分類行 → 列名行 → 薬品行 …
  列名は区分ごとに異なる（外用薬・注射薬には「標準用法用量」が無い）。
  薬品名が空の行は直前の薬品の続き:
    規格あり → 同じ薬品の別規格（薬品名・成分名・先発品を引き継ぐ）
    規格なし → 用法・備考の追記行（直前の行に追記）
"""
import argparse
import csv
import re
import unicodedata
from pathlib import Path

COLUMNS = ["No", "区分", "薬効コード", "薬効分類", "薬品名", "規格・単位",
           "標準用法用量", "成分名", "備考", "先発品", "元CSV行"]
TEXT_FIELDS = ["薬品名", "規格・単位", "標準用法用量", "成分名", "備考", "先発品"]
INHERIT = ["薬品名", "成分名", "先発品"]
PAGE_NOTES = {"（次ページへ続く）", "（前ページからの続き）", "(次ページへ続く)", "(前ページからの続き)"}
CATEGORY_RE = re.compile(r"^\s*(\d{3})\s+(\D.*)$")
# 薬品名の列に規格だけが入った行（ミルセラのように元Excelでセル結合されていた薬品）
SPEC_ONLY_RE = re.compile(r"^[\d.]+\s*(μg|mg|g|mL|単位)/")
SPACES_RE = re.compile(r"[ 　]{2,}")


def clean(text: str) -> str:
    """NFKC正規化し、セル内の改行・連続空白（印刷時の折り返し）を改行に揃える。"""
    text = unicodedata.normalize("NFKC", text)
    lines = []
    for part in SPACES_RE.sub("\n", text).split("\n"):
        part = part.strip()
        if not part:
            continue
        # 「50mg/」「…ARNI)/」のように / で折り返した行は直結する
        if lines and lines[-1].endswith("/"):
            lines[-1] += part
        else:
            lines.append(part)
    return "\n".join(lines)


def append(base: str, extra: str) -> str:
    return f"{base}\n{extra}" if base and extra else base or extra


def convert(path: Path) -> list[dict]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        raw = list(csv.reader(f))

    out: list[dict] = []
    orphans: list[dict] = []  # 薬品名の行がまだ来ていない規格行（ミルセラ）
    section = code = category = category_note = ""
    header: list[str] = []

    def emit(rec: dict) -> None:
        rec["備考"] = append(rec["備考"], rec.pop("_note"))
        out.append(rec)

    def flush_orphans(owner: dict | None) -> None:
        """保留中の規格行を owner（None なら直前の薬品）の別規格として確定する。"""
        base = owner or out[-1]
        base["成分名"] = base["成分名"] or next((o["成分名"] for o in orphans if o["成分名"]), "")
        for o in orphans:
            for k in INHERIT:
                o[k] = o[k] or base[k]
            emit(o)
        orphans.clear()

    for lineno, row in enumerate(raw, start=1):
        cells = [c.strip() for c in row]
        if not any(cells):
            continue
        first = cells[0]
        if first.startswith("≪"):
            section = first.strip("≪≫")
            continue
        if first == "薬品名":
            header = cells
            continue
        rest = [c for c in cells[1:] if c]
        if (first and all(c.startswith("※") for c in rest)
                and (m := CATEGORY_RE.match(unicodedata.normalize("NFKC", first)))):
            # 「811 あへんアルカロイド系麻薬,※他の麻薬の処方については…」の注記は分類内の各薬品の備考へ
            code, category = m.group(1), m.group(2).strip()
            category_note = "\n".join(clean(c) for c in rest)
            continue
        if not header:
            raise SystemExit(f"{lineno}行目: 列名行より前にデータがあります: {row}")

        rec = {k: "" for k in TEXT_FIELDS}
        for name, value in zip(header, cells):
            # 列名の無い6列目（外用・注射）は備考のはみ出し
            key = name if name in rec else "備考"
            if value in PAGE_NOTES:
                continue
            rec[key] = append(rec[key], clean(value))
        if not any(rec.values()):
            continue  # ページ跨ぎの注記だけの行
        rec = {"区分": section, "薬効コード": code, "薬効分類": category,
               **rec, "元CSV行": str(lineno), "_note": category_note}

        if not rec["規格・単位"] and SPEC_ONLY_RE.match(rec["薬品名"]):
            rec["規格・単位"], rec["薬品名"] = rec["薬品名"], ""
            orphans.append(rec)
            continue

        # 「※他の％と規格については要相談」「※必要時に購入」のように薬品名・規格に混ざった注記は備考へ
        for k in ("薬品名", "規格・単位"):
            lines = rec[k].split("\n")
            notes = [s for s in lines if s.startswith("※")]
            if notes:
                rec[k] = "\n".join(s for s in lines if not s.startswith("※"))
                rec["備考"] = append(rec["備考"], "\n".join(notes))
        # 「ミルセラ注シリンジ　…　100μg/0.3mL/シリンジ」のように規格が薬品名セルに入っている場合
        merged = not rec["規格・単位"] and "\n" in rec["薬品名"]
        if merged:
            rec["薬品名"], rec["規格・単位"] = rec["薬品名"].split("\n", 1)
        # 薬品名の改行は印刷時の折り返しなので詰める
        rec["薬品名"] = rec["薬品名"].replace("\n", "")
        if orphans:
            # 結合セルの薬品なら手前の規格行はこの薬品のもの、そうでなければ直前の薬品の続き
            flush_orphans(rec if merged else None)

        if not rec["薬品名"]:
            if not out:
                raise SystemExit(f"{lineno}行目: 先頭行の薬品名が空です")
            prev = out[-1]
            if not rec["規格・単位"]:
                # 用法・備考の追記行（フィコンパ、エンレストなど）
                for k in ("標準用法用量", "備考", "成分名", "先発品"):
                    prev[k] = append(prev[k], rec[k])
                prev["元CSV行"] += f",{lineno}"
                continue
            for k in INHERIT:
                rec[k] = rec[k] or prev[k]
        emit(rec)
    if orphans:
        flush_orphans(None)

    for i, rec in enumerate(out, start=1):
        rec["No"] = i
    return out


def write_csv(rows: list[dict], path: Path) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(rows)


def write_xlsx(rows: list[dict], path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "院内医薬品集"
    ws.append(COLUMNS)
    for rec in rows:
        ws.append([rec[c] for c in COLUMNS])

    widths = {"No": 6, "区分": 8, "薬効コード": 8, "薬効分類": 22, "薬品名": 36, "規格・単位": 18,
              "標準用法用量": 30, "成分名": 30, "備考": 30, "先発品": 28, "元CSV行": 9}
    for idx, col in enumerate(COLUMNS, start=1):
        ws.column_dimensions[ws.cell(1, idx).column_letter].width = widths[col]
    top = Alignment(vertical="top", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = top
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="DDEBF7")
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("src", type=Path, help="院内医薬品集のCSV")
    ap.add_argument("--out", type=Path, default=Path("data/formulary/院内医薬品集"),
                    help="出力先（拡張子なし。.csv と .xlsx を作る）")
    args = ap.parse_args()

    rows = convert(args.src.expanduser())
    args.out.parent.mkdir(parents=True, exist_ok=True)
    write_csv(rows, args.out.with_suffix(".csv"))
    write_xlsx(rows, args.out.with_suffix(".xlsx"))
    by_section: dict[str, int] = {}
    for r in rows:
        by_section[r["区分"]] = by_section.get(r["区分"], 0) + 1
    print(f"出力: {args.out}.csv / .xlsx（{len(rows)}行 / 薬効分類 {len({r['薬効コード'] + r['薬効分類'] for r in rows})}種）")
    print("  " + " / ".join(f"{k} {v}行" for k, v in by_section.items()))


if __name__ == "__main__":
    main()
