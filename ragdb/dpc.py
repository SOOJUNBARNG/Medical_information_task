"""DPC電子点数表（Excel）→ raw/dpc.csv への変換。

1行 = 1診断群分類番号（14桁）。診断群分類番号の構成:
  [0:2] MDC  [2:6] 分類  [6] 病態等  [7] 年齢等  [8:10] 手術
  [10] 手術・処置等1  [11] 手術・処置等2  [12] 定義副傷病  [13] 重症度等
桁の値（対応コード）を各シートで引き、手術・処置の具体名を dpc_name に含める。

列の定義:
  surgery_flag  … 手術区分の対応コード（[8:10]。99=手術なし, 97=その他の手術, xx=区分なし）
  resource_flag … 手術・処置等1/2 の対応コード（[10:12]。医療資源の投入区分）
"""
from __future__ import annotations

import csv
import unicodedata
from collections import defaultdict
from pathlib import Path

COLUMNS = [
    "mdc_code", "mdc_name", "dpc_code", "dpc_name", "surgery_flag", "resource_flag",
    "disease_name", "surgery", "procedure1", "procedure2", "comorbidity", "severity",
    "days", "points", "icd_codes",
]
MAX_NAMES = 5  # 1項目あたり本文に含める具体名の上限（512トークンに収めるため）


def _n(v) -> str:
    if v is None:
        return ""
    return unicodedata.normalize("NFKC", str(v)).replace("　", " ").strip()


def _sheet(wb, prefix: str):
    for name in wb.sheetnames:
        if _n(name).startswith(prefix):
            return wb[name]
    raise KeyError(f"シートが見つかりません: {prefix}")


def _data_rows(ws, key_col: int = 0):
    """先頭のヘッダー行を飛ばし、MDCコード（2桁数字）で始まる行だけ返す。"""
    for r in ws.iter_rows(values_only=True):
        v = _n(r[key_col]) if len(r) > key_col else ""
        if len(v) == 2 and v.isdigit():
            yield r


def _join(names: list[str]) -> str:
    uniq = list(dict.fromkeys(n for n in names if n))
    s = "、".join(uniq[:MAX_NAMES])
    return s + (f" ほか{len(uniq) - MAX_NAMES}件" if len(uniq) > MAX_NAMES else "")


def parse_xlsx(path: Path) -> list[dict]:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True)

    mdc = {_n(r[0]): _n(r[1]) for r in _data_rows(_sheet(wb, "1)"))}
    bunrui = {(_n(r[0]), _n(r[1])): _n(r[2]) for r in _data_rows(_sheet(wb, "2)"))}

    icd = defaultdict(list)
    for r in _data_rows(_sheet(wb, "4)")):
        icd[(_n(r[0]), _n(r[1]))].append(f"{_n(r[3])} {_n(r[2])}")

    # 対応コード → 具体名
    surg = defaultdict(list)
    for r in _data_rows(_sheet(wb, "6)")):
        surg[(_n(r[0]), _n(r[1]), _n(r[5]).zfill(2))].append(_n(r[6]))
    proc1, proc2, sub = defaultdict(list), defaultdict(list), defaultdict(list)
    for r in _data_rows(_sheet(wb, "7)")):
        proc1[(_n(r[0]), _n(r[1]), _n(r[2]))].append(_n(r[5]))
    for r in _data_rows(_sheet(wb, "8)")):
        proc2[(_n(r[0]), _n(r[1]), _n(r[2]))].append(_n(r[4]))
    for r in _data_rows(_sheet(wb, "9)")):
        sub[(_n(r[0]), _n(r[1]), _n(r[2]))].append(_n(r[4]))

    rows = []
    for r in _sheet(wb, "11)").iter_rows(values_only=True):
        code = _n(r[2]) if len(r) > 2 else ""
        if len(code) != 14 or not code[:2].isdigit():
            continue
        m, b = code[:2], code[2:6]
        disease, surgery, p1, p2, cc, sev = (_n(x) for x in r[3:9])
        s_code, p1_code, p2_code, cc_code = code[8:10], code[10], code[11], code[12]

        def detail(label: str, value: str, table: dict, key_code: str) -> str:
            if not value or value == "なし":
                return f"{label}なし" if value == "なし" else ""
            names = table.get((m, b, key_code), []) if "x" not in key_code else []
            return f"{label}{value}" + (f"（{_join(names)}）" if names else "")

        parts = [disease]
        s_names = surg.get((m, b, s_code), []) if s_code not in ("xx", "97", "99") else []
        if surgery:
            parts.append(f"手術:{surgery}" + (f"（{_join(s_names)}）" if s_names and surgery != "なし" else ""))
        parts += [
            detail("手術・処置等1:", p1, proc1, p1_code),
            detail("手術・処置等2:", p2, proc2, p2_code),
            detail("定義副傷病:", cc, sub, cc_code),
        ]
        if sev:
            parts.append(f"重症度等:{sev}")
        days = "/".join(_n(x) for x in r[9:12])
        points = "/".join(_n(x) for x in r[12:15])
        rows.append({
            "mdc_code": m,
            "mdc_name": mdc.get(m, ""),
            "dpc_code": code,
            "dpc_name": " ".join(p for p in parts if p),
            "surgery_flag": s_code,
            "resource_flag": code[10:12],
            "disease_name": bunrui.get((m, b), disease),
            "surgery": surgery,
            "procedure1": p1,
            "procedure2": p2,
            "comorbidity": cc,
            "severity": sev,
            "days": days if days.strip("/") else "",
            "points": points if points.strip("/") else "",
            "icd_codes": _join(icd.get((m, b), [])),
        })
    return rows


def write_csv(rows: list[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(rows)
