"""パイプラインの単体テスト（ネットワーク・Ollama不要）。"""
from __future__ import annotations

import csv
import sqlite3
import struct
from pathlib import Path

import numpy as np
import pytest

from ragdb import chunking, config, db, dpc, icd10, pmda
from ragdb.embed import embed_all, load_done_ids, pack, unpack
from ragdb.search import Index
from ragdb.util import clean, token_len

FIXTURES = Path(__file__).parent / "fixtures"


# ── 前処理 ────────────────────────────────────────────────────
def test_clean_rules():
    assert clean("全角　スペース") == "全角 スペース"
    assert clean("a  \n\n  b\tc") == "a b c"
    assert clean("<p>本文<br/>続き</p>") == "本文 続き"
    assert clean("制御\x00文字\x1f") == "制御 文字"
    assert clean(None) == ""


# ── ICD-10 ────────────────────────────────────────────────────
ESTAT_ROWS = [
    ("I", "感染症及び寄生虫症（A00－B99）"),
    ("A00-A09", "腸管感染症（A00－A09）"),
    ("A08", "ウイルス性及びその他の明示された腸管感染症"),
    ("A08.5", "その他の明示された腸管感染症"),
    ("A08.5b", "その他"),
    ("A17†", "神経系結核"),
    ("A17.0†", "結核性髄膜炎（G01*）"),
    ("II", "新生物＜腫瘍＞（C00－D48）"),
    ("C00-C97", "悪性新生物＜腫瘍＞（C00－C97）"),
    ("C00-C14", "口唇，口腔及び咽頭の悪性新生物＜腫瘍＞（C00－C14）"),
    ("C00", "口唇の悪性新生物＜腫瘍＞"),
    ("C15-C26", "消化器の悪性新生物＜腫瘍＞（C15－C26）"),
    ("C15", "食道の悪性新生物＜腫瘍＞"),
]


def test_icd10_hierarchy():
    rows = {r["code"]: r for r in icd10.to_rows(ESTAT_ROWS)}
    assert list(rows["A08.5b"].keys()) == icd10.COLUMNS
    h = rows["A08.5b"]["hierarchy"]
    assert h.startswith("A08.5 その他の明示された腸管感染症 > A08 ")
    assert "A00-A09 腸管感染症" in h and h.endswith("第I章 感染症及び寄生虫症（A00－B99）")
    # 剣印付きコードの親子
    assert rows["A17.0†"]["hierarchy"].startswith("A17† 神経系結核")
    # 入れ子の中間分類: C15-C26 は C00-C14 の下ではなく C00-C97 の下
    assert rows["C15"]["hierarchy"].startswith("C15-C26 消化器の悪性新生物")
    assert "C00-C14" not in rows["C15"]["hierarchy"]
    assert "C00-C97" in rows["C15"]["hierarchy"]


def test_icd10_chunk_body_format(tmp_path):
    path = tmp_path / "icd10.csv"
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=icd10.COLUMNS)
        w.writeheader()
        w.writerow({"code": "X01", "name_ja": "疾病　名", "name_en": "", "definition": "定義文",
                    "includes": "包含A", "excludes": "除外B", "hierarchy": "上位"})
        w.writerow({"code": "X02", "name_ja": "", "name_en": "", "definition": "", "includes": "",
                    "excludes": "", "hierarchy": ""})  # 空はスキップ
    chunks = list(chunking.icd10_chunks(path))
    assert len(chunks) == 1
    assert chunks[0]["body"] == "疾病 名\n定義文\n\n分類: 上位\n\n包含: 包含A\n\n除外: 除外B"
    assert chunks[0]["category"] == "icd10" and chunks[0]["code"] == "X01"


# ── DPC ───────────────────────────────────────────────────────
def _dpc_workbook(path: Path) -> None:
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    hdr = [["h"], ["h"]]
    sheets = {
        "１）ＭＤＣ名称": [["01", "神経系疾患"]],
        "２）分類名称": [["01", "0010", "脳腫瘍"]],
        "４）ＩＣＤ": [["01", "0010", "脳の悪性新生物＜腫瘍＞", "C71$"]],
        "６）手術 ": [["01", "0010", None, "04", None, "03", "頭蓋内腫瘍摘出術　その他のもの", "K1692"]],
        "７）手術・処置等１": [["01", "0010", "1", "1", None, "造影剤注入手技", "E0033"]],
        "８）手術・処置等２": [["01", "0010", "1", "4", "中心静脈注射", "G005"]],
        "９）定義副傷病名": [["01", "0010", "1", "3", "水痘肺炎", "B012"]],
    }
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in hdr + rows:
            ws.append(r)
    ws = wb.create_sheet("11）診断群分類点数表")
    for r in [[None], [None], [None, "番号"], [None]]:
        ws.append(r)
    ws.append([None, 1, "010010xx03x10x", "脳腫瘍", "頭蓋内腫瘍摘出術等", "", "１あり", "なし", "",
               13, 27, 90, 3404, 2290, 1947, "0"])
    ws.append([None, 2, "01021xxxxx00xx", "認知症", "", "なし", "なし", "", "", "", "", "", "", "", "", "0"])
    wb.save(path)


def test_dpc_parse_and_chunk(tmp_path):
    xlsx = tmp_path / "dpc.xlsx"
    _dpc_workbook(xlsx)
    rows = dpc.parse_xlsx(xlsx)
    assert [r["dpc_code"] for r in rows] == ["010010xx03x10x", "01021xxxxx00xx"]
    r = rows[0]
    assert r["mdc_name"] == "神経系疾患" and r["surgery_flag"] == "03" and r["resource_flag"] == "x1"
    assert "頭蓋内腫瘍摘出術 その他のもの" in r["dpc_name"]
    assert "手術・処置等2:1あり（中心静脈注射）" in r["dpc_name"]
    assert r["days"] == "13/27/90" and r["points"] == "3404/2290/1947"

    csv_path = tmp_path / "dpc.csv"
    dpc.write_csv(rows, csv_path)
    chunks = list(chunking.dpc_chunks(csv_path))
    assert chunks[0]["body"].startswith("MDC01 神経系疾患 / 010010xx03x10x 脳腫瘍 ")
    assert "対象ICD-10: C71$ 脳の悪性新生物<腫瘍>" in chunks[0]["body"]
    assert all(token_len(c["body"]) <= config.MAX_TOKENS for c in chunks)


# ── PMDA ──────────────────────────────────────────────────────
def test_pmda_collect_and_sections(tmp_path):
    src = tmp_path / "src"
    src.mkdir()
    long_text = "".join(f"重大な副作用{i}として出血（頻度不明）があらわれることがあるので観察を十分に行うこと。" for i in range(120))
    xml = (FIXTURES / "pmda" / "sample_insert.xml").read_text(encoding="utf-8").replace("REPEAT_ADVERSE", long_text)
    (src / "a.xml").write_text(xml, encoding="utf-8")
    out = tmp_path / "pmda"
    n, skipped = pmda.collect(src, out)
    assert n == 1 and not skipped
    assert (out / "22100AMX00001.xml").exists()

    doc = pmda.parse_sections(out / "22100AMX00001.xml")
    assert doc["title"] == "テストアスピリン錠100mg"
    sections = dict(doc["sections"])
    assert set(sections) == {"警告", "禁忌", "効能又は効果", "用法及び用量", "副作用"}
    assert "併用禁忌" not in sections["禁忌"]
    assert "Test Aspirin" not in doc["title"]

    chunks = list(chunking.pmda_chunks(out))
    by_section = {}
    for c in chunks:
        by_section.setdefault(c["section"], []).append(c)
    assert "警告" not in by_section  # 50文字未満はスキップ
    assert by_section["用法及び用量"][0]["body"].startswith("通常、成人にはアスピリンとして100mg")
    assert "テストアスピリン錠100mg" in by_section["用法及び用量"][0]["embed_text"]
    adverse = by_section["副作用"]
    assert len(adverse) > 1
    assert all(token_len(c["body"]) <= config.MAX_TOKENS for c in adverse)
    assert len({c["id"] for c in chunks}) == len(chunks)


def test_split_overlap_shares_64_tokens():
    from ragdb.util import _encoder

    enc = _encoder()
    text = "".join(f"これは{i}番目の文で、薬物動態に関する説明が続きます。" for i in range(200))
    parts = chunking.split_with_overlap(text, 512, 64)
    assert len(parts) >= 3
    for a, b in zip(parts, parts[1:]):
        tail = enc.decode(enc.encode(a)[-64:]).lstrip("�").strip()
        assert b.startswith(tail[-20:]) or tail[-20:] in b[:200]
    assert all(token_len(p) <= 512 for p in parts)


# ── 埋め込み（進捗・再開）と DB 構築 ───────────────────────────
class FakeEmbedder:
    """文字バイグラムのハッシュで決定的なベクトルを作る（Ollama不要）。"""

    model = "fake"
    dim = 768

    def __init__(self, fail_after: int | None = None):
        self.calls = 0
        self.fail_after = fail_after

    def _vec(self, text: str) -> bytes:
        v = np.zeros(self.dim, dtype=np.float32)
        for a, b in zip(text, text[1:]):
            v[hash(a + b) % self.dim] += 1.0
        return pack(v.tolist())

    def embed_document(self, text: str) -> bytes:
        self.calls += 1
        if self.fail_after is not None and self.calls > self.fail_after:
            raise KeyboardInterrupt
        return self._vec(text)

    def embed_query(self, text: str, prefix: str = "") -> bytes:
        return self._vec(text)


def _chunks(n: int) -> list[dict]:
    names = ["急性心筋梗塞", "肺炎", "腰痛症", "脳梗塞", "糖尿病"]
    return [chunking._make("icd10", f"T{i:03d}", names[i % 5], None, f"{names[i % 5]} の説明 {i}") for i in range(n)]


def test_embed_resume_after_interrupt(tmp_path):
    progress, store = tmp_path / "progress.jsonl", tmp_path / "emb.sqlite"
    chunks = _chunks(50)
    with pytest.raises(KeyboardInterrupt):
        embed_all(chunks, FakeEmbedder(fail_after=23), progress=progress, store=store,
                  parallel=1, every=10, show_progress=False)
    done = load_done_ids(progress, "fake")
    assert 0 < len(done) < 50
    # 進捗に載っているIDは必ず保存済み
    con = sqlite3.connect(store)
    saved = {r[0] for r in con.execute("SELECT chunk_id FROM embeddings")}
    assert done <= saved

    emb = FakeEmbedder()
    stats = embed_all(chunks, emb, progress=progress, store=store, parallel=4, every=10, show_progress=False)
    assert stats["skipped"] == len(done) and emb.calls == 50 - len(done)
    assert load_done_ids(progress, "fake") == {c["id"] for c in chunks}
    # 別モデルの進捗は流用しない
    assert load_done_ids(progress, "nomic-embed-text") == set()


def test_build_db_and_search(tmp_path):
    chunks = _chunks(30)
    emb = FakeEmbedder()
    embeddings = {c["id"]: emb.embed_document(c["body"]) for c in chunks}
    out = tmp_path / "rag.db"
    counts = db.build(out, chunks, embeddings, model="fake", batch_size=7)
    assert counts == {"icd10": 30}
    assert not Path(f"{out}-wal").exists()

    con = sqlite3.connect(out)
    assert con.execute("PRAGMA journal_mode").fetchone()[0] == "delete"
    assert con.execute("SELECT COUNT(*) FROM chunks WHERE length(embedding) != 3072").fetchone()[0] == 0
    blob = con.execute("SELECT embedding FROM chunks LIMIT 1").fetchone()[0]
    assert len(struct.unpack("768f", blob)) == 768 and unpack(blob) == struct.unpack("<768f", blob)
    meta = db.read_meta(con)
    assert meta["embed_model"] == "fake" and meta["embed_dim"] == "768"

    with pytest.raises(FileExistsError):
        db.build(out, chunks, embeddings)

    idx = Index(out)
    hits, _ = idx.search(emb, "急性心筋梗塞", top_k=3)
    assert hits[0].title == "急性心筋梗塞" and hits[0].score > 0.5
    assert all(h.category == "icd10" for h in idx.search(emb, "肺炎", category="icd10")[0])
    assert idx.search(emb, "肺炎", category="dpc")[0] == []


def test_build_requires_all_embeddings(tmp_path):
    chunks = _chunks(3)
    with pytest.raises(KeyError):
        db.build(tmp_path / "rag.db", chunks, {})


def test_clean_keeps_medical_angle_brackets():
    assert clean("悪性新生物<腫瘍>") == "悪性新生物<腫瘍>"
    assert clean("ジカ<Zika>ウイルス病") == "ジカ<Zika>ウイルス病"
    assert clean('<span class="x">本文</span>') == "本文"


# ── 院内医薬品集 ──────────────────────────────────────────────
def test_formulary_rows_and_chunks(tmp_path):
    from ragdb import formulary

    path = tmp_path / "院内医薬品集.csv"
    path.write_text(
        "No,区分,薬効コード,薬効分類,薬品名,規格・単位,標準用法用量,成分名,備考,先発品,元CSV行\n"
        '1,内用薬,113,抗てんかん薬,アレビアチン散10%,100mg/g,1日200~300mg 分3,フェニトイン,TDM対象医薬品,,18\n'
        "2,内用薬,113,抗てんかん薬,アレビアチン錠100mg,100mg,,,,,19\n"
        '3,内用薬,112,催眠鎮静剤,ジアゼパム錠2「サワイ」,2mg,1日4~40mg,ジアゼパム,,"2mgセルシン錠\nホリゾン錠2mg",7\n'
        "4,外用薬,264,消炎剤  (次ページへ続く),デルモゾールGローション,10mL/本,,ワルファリンカリウム,,,8\n",
        encoding="utf-8-sig",
    )
    rows = formulary.read(path)
    assert rows[1]["成分名"] == "フェニトイン"  # 同じ薬剤の別規格は成分名を引き継ぐ
    assert rows[3]["薬効分類"] == "消炎剤"
    assert formulary.keywords(rows[2]) == {"ジアゼパム", "セルシン", "ホリゾン"}
    assert "ワルファリン" in formulary.keywords(rows[3])

    chunks = list(chunking.formulary_chunks(path))
    assert [c["category"] for c in chunks] == ["formulary"] * 4
    assert chunks[1]["body"] == ("院内採用薬: アレビアチン錠100mg（内用薬） / 規格: 100mg / 成分名: フェニトイン"
                                 " / 薬効分類: 113 抗てんかん薬 / 標準用法用量: 記載なし")
    assert "先発品: 2mgセルシン錠、ホリゾン錠2mg" in chunks[2]["body"]

    emb = FakeEmbedder()
    out = tmp_path / "rag.db"
    db.build(out, chunks, {c["id"]: emb.embed_document(c["body"]) for c in chunks}, model="fake")
    idx = Index(out)
    hits, _ = idx.search(emb, "ホリゾンの院内採用は？", top_k=3)
    assert hits[0].title == "ジアゼパム錠2「サワイ」" and hits[0].score == 1.0
    assert len({h.id for h in hits}) == len(hits)
